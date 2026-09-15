"""Кабінет викладача: журнал оцінок (CSV) і сигнали аномалій.

Телеметрія працює на політику, а не на звинувачення (п. 3.4, 6.4): тут лише
показуємо СИРІ сигнали (вставки, перемикання вкладки, час на пункт). Пороги
автоворідингу підбираються за даними пілоту — навмисно НЕ автоматизовано тут.
"""

from __future__ import annotations

import csv
import io

from sqlalchemy import select
from sqlalchemy.orm import Session

import engine
from app import models
from app.config import SECRET
from app.service import ServiceError

RESULT_COLS = [
    "group", "full_name", "test_key", "attempt_no", "status",
    "score", "max_score", "percent", "started_at", "finished_at",
]


def _test_or_404(db: Session, test_key: str) -> models.Test:
    test = db.scalar(select(models.Test).where(models.Test.key == test_key))
    if test is None:
        raise ServiceError(404, "тест не знайдено")
    return test


def _attempts(db: Session, test: models.Test) -> list[models.Attempt]:
    return list(
        db.scalars(select(models.Attempt).where(models.Attempt.test_id == test.id))
    )


def _dt(value) -> str:
    return value.isoformat(sep=" ", timespec="seconds") if value else ""


def result_rows(db: Session, test_key: str) -> list[dict]:
    test = _test_or_404(db, test_key)
    rows = []
    for a in _attempts(db, test):
        pct = round(100 * a.score / a.max_score) if a.max_score else 0
        rows.append(
            {
                "group": a.student.group_name,
                "full_name": a.student.full_name,
                "test_key": test.key,
                "attempt_no": a.attempt_no,
                "status": a.status,
                "score": a.score,
                "max_score": a.max_score,
                "percent": pct,
                "started_at": _dt(a.started_at),
                "finished_at": _dt(a.finished_at),
            }
        )
    rows.sort(key=lambda r: (r["group"], r["full_name"], r["attempt_no"]))
    return rows


def results_csv(db: Session, test_key: str) -> str:
    rows = result_rows(db, test_key)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=RESULT_COLS)
    writer.writeheader()
    writer.writerows(rows)
    # BOM, щоб Excel правильно показав кирилицю.
    return "﻿" + buf.getvalue()


# Українські заголовки для Excel-журналу (у тому ж порядку, що RESULT_COLS).
_XLSX_HEADERS = [
    "Група", "Прізвище та ім'я", "Тест", "Спроба", "Статус",
    "Бал", "Максимум", "Відсоток", "Розпочато", "Завершено",
]


def results_xlsx(db: Session, test_key: str) -> bytes:
    """Журнал оцінок як справжній .xlsx — відкривається в Excel одразу в колонки
    (без проблем із роздільником, на відміну від CSV в українській локалі)."""
    from openpyxl import Workbook  # ліниво: залежність потрібна лише тут
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    rows = result_rows(db, test_key)
    wb = Workbook()
    ws = wb.active
    ws.title = "Журнал"
    ws.append(_XLSX_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r[k] for k in RESULT_COLS])
    ws.freeze_panes = "A2"  # шапка лишається зверху при прокрутці
    for i, width in enumerate([14, 26, 14, 8, 10, 8, 10, 10, 20, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def anomalies(db: Session, test_key: str) -> list[dict]:
    test = _test_or_404(db, test_key)
    out = []
    for a in _attempts(db, test):
        events = list(
            db.scalars(select(models.Event).where(models.Event.attempt_id == a.id))
        )
        paste = sum(1 for e in events if e.type == "paste")
        blur = sum(1 for e in events if e.type == "blur")
        first_ms = [
            e.payload["ms"]
            for e in events
            if e.type == "first_input" and isinstance(e.payload, dict) and "ms" in e.payload
        ]
        secs = [
            (q.submitted_at - q.started_at).total_seconds()
            for q in a.questions
            if q.started_at and q.submitted_at
        ]
        flags = []
        if paste:
            flags.append("вставка")
        if blur:
            flags.append("перемикання вкладки")
        out.append(
            {
                "attempt_id": a.id,
                "group": a.student.group_name,
                "full_name": a.student.full_name,
                "attempt_no": a.attempt_no,
                "status": a.status,
                "score": a.score,
                "max_score": a.max_score,
                "finished_at": _dt(a.finished_at),
                "paste_count": paste,
                "blur_count": blur,
                "min_seconds_per_question": round(min(secs), 1) if secs else None,
                "avg_first_input_ms": round(sum(first_ms) / len(first_ms)) if first_ms else None,
                "flags": flags,
            }
        )
    # Спершу спроби з прапорцями — їх викладач перегляне насамперед.
    out.sort(key=lambda r: (not r["flags"], r["group"], r["full_name"]))
    return out


def attempt_detail(db: Session, attempt_id: int) -> dict:
    """Поіменний розбір спроби: по кожному питанню — що студент увів, який
    еталон і скільки балів. Еталони перераховуються з seed (у базі їх немає).
    Лише для викладача (ендпоінт за токеном)."""
    attempt = db.get(models.Attempt, attempt_id)
    if attempt is None:
        raise ServiceError(404, "спробу не знайдено")

    questions = []
    for aq in sorted(attempt.questions, key=lambda q: q.ordinal):
        q = engine.build(
            aq.question_key,
            SECRET,
            attempt.student.ident,
            attempt.test.key,
            attempt.attempt_no,
            aq.reissue,
        )
        stored = {a.part_key: a for a in aq.answers}
        parts = []
        for p in q.parts:
            a = stored.get(p.key)
            got = a.score if a else 0.0
            mx = a.max_points if a else p.points
            parts.append(
                {
                    "label": p.label,
                    "raw": a.raw if a else "",
                    "expected": str(p.answer),
                    "correct": bool(mx and got >= mx),
                    "score": got,
                    "max": mx,
                }
            )
        if aq.score is None:
            verdict = "—"
        elif aq.max_score and aq.score >= aq.max_score:
            verdict = "повністю"
        elif aq.score > 0:
            verdict = "частково"
        else:
            verdict = "ні"
        questions.append(
            {
                "ordinal": aq.ordinal,
                "question_key": aq.question_key,
                "reissue": aq.reissue,
                "score": aq.score,
                "max_score": aq.max_score,
                "verdict": verdict,
                "parts": parts,
            }
        )
    return {
        "attempt_id": attempt.id,
        "full_name": attempt.student.full_name,
        "group": attempt.student.group_name,
        "status": attempt.status,
        "score": attempt.score,
        "max_score": attempt.max_score,
        "questions": questions,
    }
